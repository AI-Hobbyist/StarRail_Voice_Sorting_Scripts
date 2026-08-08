"""Generate the per-speaker voice statistics XLSX from a sorted dataset."""
from __future__ import annotations

import argparse
import wave
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
LANG = {"CHS": "中文 - Chinese", "JP": "日语 - Japanese", "EN": "英语 - English", "KR": "韩语 - Korean"}
SHEET = {"CHS": "中文", "JP": "日语", "EN": "英语", "KR": "韩语"}
CATS = {"战斗语音 - Battle": 8, "怪物语音 - Monster": 11, "其它语音 - Others": 14, "带变量语音 - Placeholder": 17, "多人对话语音 - Conversation": 20}


def duration(path: Path) -> float:
    try:
        with wave.open(str(path), "rb") as f:
            return f.getnframes() / f.getframerate() / 86400
    except (wave.Error, EOFError):
        return 0.0


def stat_rows(dataset: Path, language: str) -> list[list[Any]]:
    root = dataset / LANG[language]
    rows = []
    if not root.is_dir():
        return rows
    for speaker in sorted((p for p in root.iterdir() if p.is_dir())):
        wavs = list(speaker.rglob("*.wav"))
        labs = list(speaker.rglob("*.lab"))
        row = [speaker.name, 0, 0, 0.0, len(wavs), len(labs), sum(map(duration, wavs)), "|"] + [0] * 15
        plain = list(speaker.glob("*.wav"))
        plainlab = list(speaker.glob("*.lab"))
        row[1:4] = [len(plain), len(plainlab), sum(map(duration, plain))]
        for folder, index in CATS.items():
            path = speaker / folder
            ws = list(path.rglob("*.wav")) if path.is_dir() else []
            ls = list(path.rglob("*.lab")) if path.is_dir() else []
            row[index:index + 3] = [len(ws), len(ls), sum(map(duration, ws))]
        rows.append(row)
    return rows


def statistics(template: Path, dataset: Path, language: str, game: str, version: str, outdir: Path, dry: bool) -> Path:
    output = outdir / f"{game}_{version}_数据统计.xlsx"
    if dry:
        return output
    book = load_workbook(template)
    ws = book[SHEET[language]]
    table = next(iter(ws.tables.values()))
    rows = stat_rows(dataset, language)
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    table.ref = f"A1:W{len(rows) + 1}"
    outdir.mkdir(parents=True, exist_ok=True)
    book.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="从已整理的数据集生成角色语音统计 XLSX")
    parser.add_argument("-l", "--language", required=True, choices=LANG, help="语言")
    parser.add_argument("--dataset", type=Path, default=ROOT / "Data/second_sorted", help="最终分类结果目录")
    parser.add_argument("--template", type=Path, default=ROOT / "数据统计_模板.xlsx", help="统计模板 XLSX")
    parser.add_argument("--game", default="星穹铁道", help="游戏名，用于输出文件名")
    parser.add_argument("--version", default="4.4", help="版本号，用于输出文件名")
    parser.add_argument("--stats-dir", type=Path, default=ROOT, help="统计 XLSX 输出目录")
    parser.add_argument("--dry-run", action="store_true", help="仅打印输出路径，不写入")
    args = parser.parse_args()
    output = statistics(
        args.template,
        args.dataset,
        args.language.upper(),
        args.game,
        args.version,
        args.stats_dir,
        args.dry_run,
    )
    print(output)


if __name__ == "__main__":
    main()
